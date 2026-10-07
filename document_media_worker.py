"""Fixed local renderer. Inputs are data, never commands, URLs or executable code."""
import io
import json
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request
import wave
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont


def _rgb(value):
    from pptx.dml.color import RGBColor
    value=value.lstrip('#')
    return RGBColor(int(value[0:2],16),int(value[2:4],16),int(value[4:6],16))


def _lum(value):
    value=value.lstrip('#')
    def channel(c):
        c=int(c,16)/255
        return c/12.92 if c<=0.03928 else ((c+0.055)/1.055)**2.4
    r,g,b=(channel(value[i:i+2]) for i in (0,2,4))
    return 0.2126*r+0.7152*g+0.0722*b


def _sparse_lines(slide):
    """Short body-only content (1-4 lines of up to 40 chars, no table/diagram/asset) -> its lines, else None.
    Such a page is laid out as a centred statement (1 line) or a row of numbered cards (2-4 lines)
    instead of a few lines at the top of an empty page."""
    if slide.get('table') or slide.get('diagram') or slide.get('asset_path') or (slide.get('layout') or 'content')!='content':
        return None
    lines=[l.strip().lstrip('・-*•●◦ ').strip() for l in (slide.get('body') or '').splitlines() if l.strip()]
    if not 1<=len(lines)<=4 or any(len(l)>40 for l in lines):
        return None
    return lines


def _card_size(points,width_pt):
    """Largest size (24-14pt) at which every point fits on one line, or wraps to two lines whose second
    line has at least 3 characters (no single-character orphans). CJK glyphs are ~1.05em wide."""
    for size in range(24,13,-1):
        per=int(width_pt//(size*1.05))
        if per>=4 and all(len(p)<=per or (len(p)<=2*per and len(p)-per>=3) for p in points):
            return size
    return 14


def _wrapped(text,width=48):
    lines=[]
    for line in (text or '').splitlines():
        if not line:
            lines.append('');continue
        while line:
            lines.append(line[:width]);line=line[width:]
    return lines


def _fallback_body(slide):
    """Text form of table/diagram for PDF/MP4 frames (PowerPoint draws them as real objects)."""
    body=slide.get('body','')
    if slide.get('table'):
        body+=('\n' if body else '')+'\n'.join(' ｜ '.join('' if v is None else str(v) for v in row) for row in slide['table'])
    if slide.get('diagram'):
        body+=('\n' if body else '')+' → '.join(slide['diagram']['nodes'])
    return body


def build_pptx(spec,slides,target):
    """Designed, editable PPTX: theme colours, cover/section/content layouts, real tables, box-and-arrow diagrams."""
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    from pptx.oxml.ns import qn
    # A learned design profile (design_profile.slide_palette) overrides the default palette and fonts.
    design=spec.get('design') or {}
    accent=_rgb(design.get('accent') or (spec.get('theme') or {}).get('accent','#176b5b'))
    dark,text,bg,light,muted=(_rgb(design.get(k,d)) for k,d in (('dark','#143d38'),('text','#203c39'),('bg','#f5f8f7'),('light','#e3efec'),('muted','#6b807c')))
    white=_rgb(design.get('on_accent','#ffffff'))
    heading_font,body_font=design.get('heading_font','Meiryo'),design.get('body_font','Meiryo')
    # Learned spot colour (design_profile.slide_palette 'secondary'); without it the layout is unchanged.
    spot=_rgb(design['secondary']) if design.get('secondary') else None
    def visible_on_accent(colour):
        a,b=(_lum(design.get('accent','#176b5b')),_lum(colour))
        return (max(a,b)+0.05)/(min(a,b)+0.05)>=1.5
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    doc_title=(spec.get('document_title') or slides[0]['title']).replace('\n',' ')
    def rect(page,x,y,w,h,fill,shape=MSO_SHAPE.RECTANGLE,line=None):
        s=page.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
        s.fill.solid();s.fill.fore_color.rgb=fill;s.shadow.inherit=False
        if line is None:s.line.fill.background()
        else:s.line.color.rgb=line
        return s
    def write(frame,lines,size,color,bold=False,align=None):
        frame.word_wrap=True
        for i,line in enumerate(lines or ['']):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            font=heading_font if bold else body_font
            p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.color.rgb=color;p.font.bold=bold;p.space_after=Pt(4)
            for run in p.runs:
                # East Asian typeface too, so Japanese text follows the design font.
                rpr=run._r.get_or_add_rPr();ea=rpr.find(qn('a:ea'))
                if ea is None:ea=rpr.makeelement(qn('a:ea'),{});rpr.append(ea)
                ea.set('typeface',font)
            if align is not None:p.alignment=align
    def textbox(page,x,y,w,h,lines,size,color,bold=False,align=None,anchor=None):
        box=page.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        if anchor is not None:box.text_frame.vertical_anchor=anchor
        write(box.text_frame,lines,size,color,bold,align)
        return box
    section_no=0
    for index,slide in enumerate(slides):
        layout=slide.get('layout') or ('content' if any(slide.get(k) for k in ('table','diagram','asset_path')) else 'cover' if index==0 else 'content')
        body,table,diagram,asset=slide.get('body',''),slide.get('table'),slide.get('diagram'),slide.get('asset_path')
        if layout in ('cover','section'):
            if table or diagram or asset:
                raise ValueError(f'{index+1}枚目は表・図・画像を含むためlayout=contentを指定してください。素材を省略せず別ページに分けてください。')
            if len(_wrapped(body))>7 or len(slide['title'].splitlines())>3:
                raise ValueError(f'{index+1}枚目の表紙・章扉の文字が収まりません。本文をcontentページへ分けてください。')
        if layout=='content':
            diagram_height=(0.6*len(diagram['nodes'])+0.25*(len(diagram['nodes'])-1)) if diagram and (diagram.get('type')=='stack' or diagram.get('direction')=='vertical') else 1.8 if diagram else 0
            need=(0.42*len(_wrapped(body))+0.2 if body else 0)+(0.42*len(table)+0.1 if table else 0)+diagram_height+(4.8 if asset else 0)
            if need>6.7:
                raise ValueError(f"{index+1}枚目「{slide['title'][:20]}」は本文・表・図の量がページに収まりません（本文は1行48文字で約15行、表は約14行、図・画像は本文を短くして併用）。このページを複数のページに分けてください。")
        page=prs.slides.add_slide(prs.slide_layouts[6])
        if layout=='cover':
            rect(page,0,0,16,9,accent);rect(page,0.8,5.15,6,0.08,spot if spot is not None and visible_on_accent(design['secondary']) else white)
            textbox(page,0.8,1.8,14.4,3.2,slide['title'].splitlines(),44,white,True,anchor=MSO_ANCHOR.BOTTOM)
            textbox(page,0.8,5.45,14.4,3,body.splitlines(),20,white)
        elif layout=='section':
            rect(page,0,0,16,9,bg);rect(page,0,0,5.2,9,accent)
            # Chapter ordinal (1st section = 01), not the slide number.
            section_no+=1
            textbox(page,0.7,3.4,4,1.6,[f'{section_no:02d}'],60,white,True)
            if spot is not None and visible_on_accent(design['secondary']):rect(page,0.8,5.0,1.3,0.07,spot)
            textbox(page,5.9,2.9,9.4,2,[slide['title']],36,dark,True,anchor=MSO_ANCHOR.MIDDLE)
            if body:textbox(page,5.9,5.0,9.4,3,body.splitlines(),20,text)
        else:
            rect(page,0,0,16,9,bg);rect(page,0,0,16,0.18,accent);rect(page,0.8,0.55,0.12,0.75,accent)
            if spot is not None:rect(page,1.1,1.36,1.1,0.05,spot)
            textbox(page,1.1,0.45,14,0.95,[slide['title']],28,dark,True,anchor=MSO_ANCHOR.MIDDLE)
            y=1.6
            if asset:
                with Image.open(asset) as im:aw,ah=im.size
                scale=min(14.2/aw,4.6/ah)
                page.shapes.add_picture(asset,Inches((16-aw*scale)/2),Inches(y),width=Inches(aw*scale),height=Inches(ah*scale))
                y+=ah*scale+0.2
            sparse=_sparse_lines(slide) if index>0 else None
            if sparse and len(sparse)==1:
                # One statement: large, centred in the content area, with a short rule above it.
                rect(page,0.9,3.55,1.1,0.06,spot if spot is not None else accent)
                textbox(page,0.9,3.8,14.2,1.8,sparse,32,dark,True,anchor=MSO_ANCHOR.TOP)
                body=''
            elif sparse:
                # 2-4 short points: a row of numbered cards, vertically centred.
                n=len(sparse);gap=0.4;w=(14.2-gap*(n-1))/n;h=2.7;top=1.6+(6.6-h)/2
                # One size for all cards: the longest point fits on one line, or splits into two even lines.
                card_size=_card_size(sparse,(w-0.7)*72-18)
                for i,point in enumerate(sparse):
                    x=0.9+i*(w+gap)
                    rect(page,x,top,w,h,_rgb('#ffffff'),MSO_SHAPE.RECTANGLE,light)
                    rect(page,x,top,w,0.08,spot if spot is not None else accent)
                    textbox(page,x+0.35,top+0.35,w-0.7,0.8,[f'{i+1:02d}'],28,accent,True)
                    textbox(page,x+0.35,top+1.2,w-0.7,h-1.4,[point],card_size,dark)
                body=''
            if body:
                h=0.42*len(_wrapped(body))+0.2 if (table or diagram or asset) else 6.6
                textbox(page,0.9,y,14.2,h,body.splitlines(),18 if (table or diagram or asset) else 20,text)
                y+=h+0.1
            if table:
                rows,cols=len(table),max(len(r) for r in table)
                height=min(0.42*rows,8.25-y)
                grid=page.shapes.add_table(rows,cols,Inches(0.9),Inches(y),Inches(14.2),Inches(height)).table
                for r,row in enumerate(table):
                    for c in range(cols):
                        cell=grid.cell(r,c);value=row[c] if c<len(row) else ''
                        cell.fill.solid();cell.fill.fore_color.rgb=accent if r==0 else (_rgb('#ffffff') if r%2 else light)
                        cell.text_frame.text='';write(cell.text_frame,['' if value is None else str(value)],14 if rows<=8 else 12,white if r==0 else text,r==0)
                y+=height+0.15
            if diagram:
                nodes=diagram['nodes'];n=len(nodes);space=max(1.2,8.25-y)
                if diagram.get('type')=='stack' or diagram.get('direction')=='vertical':
                    h=min(0.8,(space-0.25*(n-1))/n)
                    for i,label in enumerate(nodes):
                        top=y+i*(h+0.25)
                        s=rect(page,3,top,10,h,light,MSO_SHAPE.ROUNDED_RECTANGLE,accent)
                        s.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE;write(s.text_frame,[label],16,dark,True,PP_ALIGN.CENTER)
                        if i<n-1:rect(page,7.8,top+h+0.03,0.4,0.19,accent,MSO_SHAPE.DOWN_ARROW)
                else:
                    gap=0.55;w=(14.2-gap*(n-1))/n;h=min(1.6,space);top=y+(space-h)/2
                    for i,label in enumerate(nodes):
                        x=0.9+i*(w+gap)
                        s=rect(page,x,top,w,h,light,MSO_SHAPE.ROUNDED_RECTANGLE,accent)
                        s.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE;write(s.text_frame,[label],15,dark,True,PP_ALIGN.CENTER)
                        if i<n-1:rect(page,x+w+0.08,top+h/2-0.17,gap-0.16,0.34,accent,MSO_SHAPE.RIGHT_ARROW)
        if layout!='cover':
            left=5.6 if layout=='section' else 0.8
            rect(page,left,8.45,15.2-left,0.02,light)
            textbox(page,left,8.5,11.9-left,0.4,[doc_title[:70]],11,muted)
            textbox(page,12.7,8.5,2.5,0.4,[f'{index+1} / {len(slides)}'],11,muted,align=PP_ALIGN.RIGHT)
        if slide.get('notes'):page.notes_slide.notes_text_frame.text=slide['notes']
    prs.save(target)


def render(spec, target, ffmpeg):
    slides=spec['slides']
    design=spec.get('design') or {}
    font=Path(design.get('body_file') or 'C:/Windows/Fonts/meiryo.ttc')
    heading_file=Path(design.get('heading_file') or font)
    if not font.is_file(): font=Path('C:/Windows/Fonts/meiryo.ttc')
    if not heading_file.is_file(): heading_file=font
    if not font.is_file(): raise ValueError('日本語フォントが見つかりません。')
    bg,accent,title_color,body_color=(design.get(k,d) for k,d in (('bg','#f5f8f7'),('accent','#176b5b'),('dark','#143d38'),('text','#203c39')))
    if target.suffix=='.pptx':
        build_pptx(spec,slides,target)
        return
    slides=[dict(s,body=_fallback_body(s)) for s in slides]
    title_font=ImageFont.truetype(str(heading_file),52)
    body_font=ImageFont.truetype(str(font),32)
    with tempfile.TemporaryDirectory(prefix='saikuru-media-',dir=target.parent) as folder:
        temp=Path(folder); images=[]
        for index,slide in enumerate(slides):
            image=Image.new('RGB',(1920,1080),bg);draw=ImageDraw.Draw(image)
            draw.rectangle((0,0,1920,24),fill=accent)
            if design.get('secondary'):
                title_lines=max(1,(len(slide['title'])+29)//30)
                draw.rectangle((100,80+title_lines*65+6,280,80+title_lines*65+14),fill=design['secondary'])
            for line_index,offset in enumerate(range(0,len(slide['title']),30)):
                draw.text((100,80+line_index*65),slide['title'][offset:offset+30],font=title_font,fill=title_color)
            asset=slide.get('asset_path')
            if asset:
                source=Path(asset)
                if source.suffix=='.mp4':
                    if target.suffix!='.mp4':raise ValueError('操作動画素材の統合はMP4のみです。')
                    source=temp/f'asset-{index}.png'
                    subprocess.run([ffmpeg,'-hide_banner','-loglevel','error','-nostdin','-i',asset,'-frames:v','1',str(source)],check=True,timeout=30,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
                with Image.open(source) as original:
                    visual=original.convert('RGB');visual.thumbnail((1720,620));image.paste(visual,((1920-visual.width)//2,210+(620-visual.height)//2))
                draw.text((100,160),'実画面素材',font=body_font,fill=accent)
            y=850 if asset else 230
            sparse=_sparse_lines(slide) if index>0 else None
            if sparse and len(sparse)==1:
                draw.rectangle((100,470,240,478),fill=design.get('secondary') or accent)
                text_line=sparse[0]
                for offset in range(0,len(text_line),26):
                    draw.text((100,500+offset//26*70),text_line[offset:offset+26],font=title_font,fill=title_color)
                slide=dict(slide,body='')
            elif sparse:
                n=len(sparse);gap=40;w=(1720-gap*(n-1))//n;h=340;top=230+(790-h)//2
                per=max(4,(w-80)//36)
                for i,point in enumerate(sparse):
                    x=100+i*(w+gap)
                    draw.rectangle((x,top,x+w,top+h),fill='#ffffff',outline=design.get('light','#e3efec'),width=2)
                    draw.rectangle((x,top,x+w,top+10),fill=design.get('secondary') or accent)
                    draw.text((x+40,top+40),f'{i+1:02d}',font=title_font,fill=accent)
                    for row,offset in enumerate(range(0,len(point),per)):
                        draw.text((x+40,top+140+row*48),point[offset:offset+per],font=body_font,fill=title_color)
                slide=dict(slide,body='')
            for line in slide['body'].splitlines():
                while line:
                    part=line[:48];line=line[48:]
                    if y>1020:raise ValueError(f"{index+1}枚目「{slide['title'][:20]}」の本文がページに収まりません（1ページは1行48文字で約15行まで。見出し下の空行も1行と数える）。このページを複数のページに分けてください。")
                    draw.text((100,y),part,font=body_font,fill=body_color);y+=52
                y+=12
            frame=temp/f'{index:03}.png';image.save(frame);images.append(frame)
        if target.suffix=='.pdf':
            from reportlab.pdfgen import canvas
            from reportlab.lib.utils import ImageReader
            pdf=canvas.Canvas(str(target),pagesize=(960,540))
            for frame in images:pdf.drawImage(ImageReader(str(frame)),0,0,960,540);pdf.showPage()
            pdf.save()
        else:
            if not ffmpeg:raise ValueError('FFmpegが見つかりません。')
            clips=[]
            # Explicit localhost only. Do not inherit proxy settings for narration.
            opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
            for index,(slide,frame) in enumerate(zip(slides,images)):
                narration=slide.get('narration','').strip()
                wav=temp/f'{index:03}.wav'
                if narration:
                    speaker=spec.get('speaker_id')
                    if type(speaker) is not int or not 0<=speaker<=10000:raise ValueError('VOICEVOX話者IDを指定してください。')
                    query=urllib.parse.urlencode({'text':narration,'speaker':speaker})
                    req=urllib.request.Request('http://127.0.0.1:50021/audio_query?'+query,data=b'',method='POST')
                    with opener.open(req,timeout=30) as r:audio_query=r.read(1000000)
                    req=urllib.request.Request('http://127.0.0.1:50021/synthesis?speaker='+str(speaker),data=audio_query,headers={'Content-Type':'application/json'},method='POST')
                    with opener.open(req,timeout=90) as r:raw=r.read(30000001)
                    if len(raw)>30000000:raise ValueError('音声サイズの上限を超えました。')
                    wav.write_bytes(raw)
                    with wave.open(io.BytesIO(raw)) as audio:duration=audio.getnframes()/audio.getframerate()
                else:duration=slide.get('duration',8)
                if not 1<=duration<=180:raise ValueError('章の長さは1〜180秒にしてください。')
                clip=temp/f'{index:03}.mp4';clips.append(clip)
                args=[ffmpeg,'-hide_banner','-loglevel','error','-nostdin','-loop','1','-i',str(frame)]
                asset=slide.get('asset_path')
                moving=bool(asset and Path(asset).suffix=='.mp4')
                if moving:args+=['-i',asset]
                if narration:args+=['-i',str(wav)]
                if moving:
                    args+=['-filter_complex','[1:v]scale=1720:620:force_original_aspect_ratio=decrease:force_divisible_by=2,pad=1720:620:(ow-iw)/2:(oh-ih)/2:color=0xf5f8f7,tpad=stop_mode=clone:stop_duration=180[v];[0:v][v]overlay=100:210[out]','-map','[out]']
                    if narration:args+=['-map','2:a']
                args+=['-t',str(duration),'-r','24','-c:v','libx264','-preset','veryfast','-pix_fmt','yuv420p']
                if narration:args+=['-c:a','aac']
                args+=[str(clip)]
                subprocess.run(args,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=240,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            if len({bool(s.get('narration','').strip()) for s in slides})>1:raise ValueError('音声あり・なしの章を混在させないでください。')
            concat=temp/'clips.txt';concat.write_text('\n'.join("file '"+p.name+"'" for p in clips),encoding='utf-8')
            subprocess.run([ffmpeg,'-hide_banner','-loglevel','error','-nostdin','-f','concat','-safe','1','-i',str(concat),'-c','copy','-movflags','+faststart',str(target)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=90,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))


if __name__=='__main__':
    try:
        render(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')),Path(sys.argv[2]),sys.argv[3])
    except ValueError as exc:
        # Input problems the agent can fix (e.g. text overflow) are reported back; other errors stay generic.
        sys.stderr.write('SAIKURU_INPUT_ERROR: '+str(exc)[:500].replace('\n',' ')+'\n')
        sys.exit(2)
