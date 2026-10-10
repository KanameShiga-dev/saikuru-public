"""Bounded document output broker; source is always read-only."""
import hashlib
import io
import json
import os
import re
from pathlib import Path
import sys
import subprocess
import tempfile
import time
import shutil
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from consultation_read_tools import execute, TOOLS
import design_profile

DOCUMENTS = {'.md', '.html', '.txt', '.svg', '.pptx', '.pdf', '.mp4', '.docx', '.xlsx', '.json'}
BINARY = ('.pptx', '.pdf', '.mp4', '.docx', '.xlsx')

def output_path(root, name):
    relative = Path(name)
    if relative.is_absolute() or any(p in ('.', '..') or p.startswith('.') for p in relative.parts):
        raise ValueError('資料の相対パスを指定してください。')
    if relative.suffix.lower() not in DOCUMENTS:
        raise ValueError('資料はmd/html/txt/svg/pptx/pdf/mp4/docx/xlsxのみです（jsonはデザインプロファイルの読み取りだけ）。')
    path = (root / relative).resolve()
    path.relative_to(root)
    # Lower-case the extension so 'X.PDF' is handled like 'x.pdf' (Windows names are case-insensitive).
    return path.with_suffix(path.suffix.lower())

def pdf_text(path, max_pages=60, max_chars=16000):
    """Text of a PDF, page by page (read only; nothing in the PDF is executed). Without this a request such as
    "check the PDF's contents" could never be completed, and the checking role kept stopping (2026-10-08)."""
    from pypdf import PdfReader
    raw = path.read_bytes()
    if not raw.startswith(b'%PDF-'):
        raise ValueError('PDFの形式を確認できません。')
    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        return {'page_count': None, 'pages': [], 'note': 'パスワード付きPDFのため本文を読めません。'}
    count, pages, used, truncated = len(reader.pages), [], 0, False
    for index, page in enumerate(reader.pages[:max_pages]):
        text = (page.extract_text() or '').strip()
        if used + len(text) > max_chars:
            text, truncated = text[:max(0, max_chars - used)], True
        used += len(text)
        pages.append({'page': index + 1, 'text': text})
        if truncated:
            break
    note = '本文の文字だけです。レイアウト・配色・図の見た目は、受け入れ時の自動画像化で利用者が確認します。'
    # PDFs made by generate_media hold each page as a picture: read those pages with the local Windows OCR
    # (the same fixed script as the attachment check; nothing is sent outside this PC).
    empty = [p for p in pages if not p['text']][:20]
    if empty:
        try:
            for page, text in zip(empty, _ocr_pdf_pages(raw, [p['page'] for p in empty])):
                page['text'], page['ocr'] = text[:max(0, max_chars - used)], True
                used += len(page['text'])
            note += ' 文字のないページは画像から文字認識（OCR）で読み取りました。誤認識が含まれる場合があります。'
        except (OSError, ValueError, RuntimeError, subprocess.TimeoutExpired):
            note += ' 画像だけのページの文字認識（OCR）ができませんでした。'
    if truncated or count > max_pages:
        note += f' 長いため途中までです（{max_chars}文字・{max_pages}ページまで）。'
    if not any(p['text'] for p in pages):
        note += ' 文字を取り出せませんでした。'
    return {'page_count': count, 'pages': pages, 'note': note}


def _ocr_pdf_pages(raw, numbers):
    import pypdfium2 as pdfium
    folder = tempfile.TemporaryDirectory(prefix='saikuru-pdf-ocr-')
    try:
        paths, document = [], pdfium.PdfDocument(raw)
        try:
            for number in numbers:
                page = document[number - 1]
                try:
                    width, height = page.get_size()
                    bitmap = page.render(scale=min(2.5, 1800 / max(width, height, 1)))
                    try:
                        target = Path(folder.name) / f'{number}.png'
                        bitmap.to_pil().convert('RGB').save(target)
                        paths.append(str(target.resolve()))
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        finally:
            document.close()
        exe = Path(os.environ.get('SystemRoot', r'C:\Windows')) / 'System32/WindowsPowerShell/v1.0/powershell.exe'
        # Same process-only policy bypass as the attachment check; only this fixed local OCR script is run.
        done = subprocess.run([str(exe), '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File', str(Path(__file__).with_name('attachment_ocr.ps1'))],
                              input=json.dumps({'paths': paths}), capture_output=True, text=True, encoding='utf-8', timeout=120,
                              creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        result = json.loads(done.stdout)
        if done.returncode or result.get('ok') is not True or len(result.get('texts') or []) != len(paths) * 2:
            raise RuntimeError('OCR unavailable')
        # Two results per image (Japanese, then English engine); the Japanese one covers mixed text.
        # Windows OCR separates every Japanese character with a space; keep spaces only around ASCII words.
        return [re.sub(r'(?<=[^\x00-\x7f]) +(?=[^\x00-\x7f])', '', str(result['texts'][index * 2])).strip()
                for index in range(len(paths))]
    finally:
        folder.cleanup()

def _own_output(root, path, job_id):
    """True when an existing path is a deliverable this job generated and nobody has changed since.
    2026-10-08: the user allowed re-generating one's own documents; any other existing file is still never overwritten."""
    if not path.exists():
        return False
    manifest = root / ('.saikuru-output-' + str(job_id) + '.json')
    try:
        records = json.loads(manifest.read_text(encoding='utf-8')) if job_id and manifest.is_file() and not manifest.is_symlink() else []
    except ValueError:
        records = []
    if path.is_file() and not path.is_symlink() and isinstance(records, list):
        key, actual = str(path.relative_to(root)), hashlib.sha256(path.read_bytes()).hexdigest()
        if any(isinstance(r, dict) and r.get('path') == key and r.get('sha256') == actual for r in records):
            return True
    raise ValueError('既存のファイルは上書きしません（この依頼で作った資料だけ作り直せます）。新しい名前を指定してください。')


def _set_aside(root, path):
    """Keep the previous version under .saikuru-versions/ before the new one is written."""
    folder = root / '.saikuru-versions'
    if folder.is_symlink():
        raise ValueError('旧版の退避先にリンクは使えません。')
    folder.mkdir(exist_ok=True)
    stamp = time.strftime('%Y%m%d-%H%M%S')
    target, number = folder / f'{path.stem}-{stamp}{path.suffix}', 1
    while target.exists():
        target, number = folder / f'{path.stem}-{stamp}-{number}{path.suffix}', number + 1
    os.replace(path, target)
    return str(target.relative_to(root))


def document_operation(root, name, args, writable):
    path = output_path(root, args.get('path', ''))
    if name == 'read_document':
        if path.suffix in ('.docx','.xlsx'):
            return dict(office_text(path), path=str(path.relative_to(root)), bytes=path.stat().st_size,
                        sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        if path.suffix == '.pptx':
            return dict(pptx_text(path), path=str(path.relative_to(root)), bytes=path.stat().st_size,
                        sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        if path.suffix == '.pdf':
            return dict(pdf_text(path), path=str(path.relative_to(root)), bytes=path.stat().st_size,
                        sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        if path.suffix == '.mp4':
            return {'path':str(path.relative_to(root)), 'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'note':'バイナリ成果物です。内容の視聴・表示確認は別途必要です。'}
        if path.stat().st_size > 200000: raise ValueError('資料が大きすぎます。')
        data = path.read_bytes()
        return {'content': data.decode('utf-8')[:16000], 'sha256': hashlib.sha256(data).hexdigest()}
    if name != 'write_document' or not writable: raise ValueError('資料変更は禁止です。')
    if path.suffix == '.json': raise ValueError('JSONは書けません（デザインプロファイルはlearn_designが作ります）。')
    content = args.get('content')
    if path.suffix.lower() == '.html' and args.get('design_profile') and isinstance(content, str):
        # Unified design: the learned profile's CSS goes first in <head> (or at the top of a fragment).
        style = '<style>\n' + design_profile.css(_design(root, args['design_profile'])) + '</style>\n'
        at = content.lower().find('<head>')
        content = content[:at + 6] + '\n' + style + content[at + 6:] if at >= 0 else style + content
    if path.suffix in ('.docx','.xlsx'):raise ValueError('Word・Excelはgenerate_officeツールを使ってください。')
    if path.suffix in ('.pptx','.pdf','.mp4'):raise ValueError('バイナリ成果物はgenerate_mediaツールを使ってください。')
    if not isinstance(content, str) or len(content.encode('utf-8')) > 200000:
        raise ValueError('資料は200KB以内です。')
    if path.suffix.lower() in ('.html','.svg') and re.search(r'<script\b|\bon\w+\s*=|javascript:',content,re.I):
        raise ValueError('資料に実行用スクリプトは含められません。')
    manifest = root / '.saikuru-document-manifest.json'
    if manifest.is_symlink() or manifest.with_suffix('.tmp').is_symlink():
        raise ValueError('資料管理ファイルのリンクは使えません。')
    records = json.loads(manifest.read_text(encoding='utf-8')) if manifest.exists() else {}
    key = str(path.relative_to(root))
    if path.exists():
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if records.get(key) != actual or args.get('expected_sha256') != actual:
            raise ValueError('既存資料・外部変更を上書きしません。新しい名前を指定してください。')
    path.parent.mkdir(parents=True, exist_ok=True)
    # Resolve again after mkdir; reject an escaped parent.
    path = output_path(root, args['path'])
    with path.open('w' if path.exists() else 'x', encoding='utf-8') as f: f.write(content)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    records[key] = digest
    temp = manifest.with_suffix('.tmp')
    temp.write_text(json.dumps(records), encoding='utf-8'); os.replace(temp, manifest)
    return {'path': key, 'sha256': digest}

def media_environment():
    from team_document_capabilities import runtime
    python,ffmpeg=runtime()
    result={'renderer':'PowerPoint/PDF/slide MP4','ffmpeg':bool(ffmpeg),'voicevox_available':False,'speakers':[]}
    try:
        opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open('http://127.0.0.1:50021/speakers',timeout=3) as response:
            speakers=json.loads(response.read(1000000))
        result['speakers']=[{'name':s['name'],'style':style['name'],'id':style['id']} for s in speakers for style in s['styles']]
        result['voicevox_available']=True
    except (OSError,ValueError,KeyError):pass
    return result

def _design(root, name):
    """Load a design profile (*.design.json) from the output folder; validated before use."""
    path = output_path(root, name)
    if not path.name.endswith('.design.json') or not path.is_file() or path.is_symlink():
        raise ValueError('design_profile は learn_design で作った保存先内の *.design.json を相対パスで指定してください。')
    return design_profile.load(path)


def learn_design(roots, args, writable):
    """Learn colours, fonts, type sizes and margins from a reference PDF (read-only) into design/<name>.design.json/.md."""
    if not writable:
        raise ValueError('デザインの学習は制作担当（builder）が行います。')
    root = roots[0]
    pdf = _office_base(roots, str(args.get('pdf_path', '')), '.pdf')
    name = str(args.get('name') or pdf.stem)
    if not re.fullmatch(r'[A-Za-z0-9_\-\u3040-\u30ff\u4e00-\u9fff]{1,40}', name):
        raise ValueError('name は40文字以内の英数字・日本語・-・_ で指定してください。')
    target = output_path(root, 'design/' + name + '.design.json')
    summary = target.with_name(name + '.design.md')
    if target.exists() or summary.exists():
        raise ValueError('同じ名前のデザインプロファイルがあります。別の name を指定してください（既存は上書きしません）。')
    target.parent.mkdir(parents=True, exist_ok=True)
    from team_document_capabilities import runtime
    python, _ = runtime()
    with tempfile.TemporaryDirectory(prefix='saikuru-design-', dir=target.parent) as folder:
        out, md = Path(folder) / 'profile.json', Path(folder) / 'profile.md'
        result = subprocess.run([python, '-X', 'utf8', str(Path(__file__).with_name('design_profile.py')), str(pdf), str(out),
                                 '--md', str(md), '--name', name], stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=300,
                                creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        if result.returncode or not out.is_file():
            detail = [l[len('SAIKURU_INPUT_ERROR: '):] for l in result.stderr.decode('utf-8', 'replace').splitlines() if l.startswith('SAIKURU_INPUT_ERROR: ')]
            raise ValueError('PDFからデザインを学習できませんでした' + ('：' + detail[-1] if detail else '（pypdf・pypdfium2 の導入とPDFを確認してください）。'))
        profile = design_profile.load(out)
        with target.open('xb') as f:
            f.write(out.read_bytes())
        with summary.open('xb') as f:
            f.write(md.read_bytes())
    rel = str(target.relative_to(root))
    return {'design_profile': rel, 'summary': str(summary.relative_to(root)),
            'colors': {k: v for k, v in profile['colors'].items() if k != 'palette'},
            'fonts': {k: profile['fonts'][k]['family'] for k in ('heading', 'body')}, 'sizes_pt': profile['sizes_pt'],
            'line_height': profile['line_height'], 'margins_mm': profile['margins_mm'], 'page': profile['page'], 'notes': profile['notes'],
            'how_to_apply': 'generate_media（PPTX・PDF・MP4）と generate_office（Word）に design_profile="' + rel + '" を指定する。'
                            'HTML は write_document に design_profile を指定するとCSSが入る（本文は h1/h2/h3/p/table と var(--color-primary) などを使う）。'
                            'Excel は対象外。'}


def generate_media(root,args,writable,job_id):
    if not writable:raise ValueError('制作はbuilderのみです。')
    path=output_path(root,args.get('path',''))
    if path.suffix not in ('.pptx','.pdf','.mp4'):raise ValueError('PPTX・PDF・MP4の名前を指定してください。')
    replace=_own_output(root,path,job_id)
    slides=args.get('slides')
    if not isinstance(slides,list) or not 1<=len(slides)<=80:raise ValueError('ページ・章は1〜80件です。')
    design=_design(root,args['design_profile']) if args.get('design_profile') else None
    theme=args.get('theme') or {}
    if not isinstance(theme,dict) or set(theme)-{'accent'} or ('accent' in theme and not re.fullmatch(r'#[0-9A-Fa-f]{6}',str(theme['accent']))):
        raise ValueError('themeは {"accent": "#RRGGBB"} の形式です。')
    if not isinstance(args.get('document_title',''),str) or len(args.get('document_title',''))>80:raise ValueError('document_titleは80文字までです。')
    for index,slide in enumerate(slides,1):
        if not isinstance(slide,dict) or set(slide)-{'title','body','narration','duration','asset_path','scene','layout','table','diagram','notes'}:raise ValueError('制作データの形式が不正です。')
        for key,limit in [('title',50),('body',1200),('narration',1500),('notes',2000)]:
            value=slide.get(key,'')
            if not isinstance(value,str) or len(value)>limit:raise ValueError(f'{index}枚目の文字数上限を超えました（{key}は{limit}文字まで）。')
        if slide.get('layout',None) not in (None,'cover','section','content'):raise ValueError(f'{index}枚目のlayoutはcover・section・contentのいずれかです。')
        table=slide.get('table')
        if table is not None and (not isinstance(table,list) or not 1<=len(table)<=20 or any(not isinstance(r,list) or not 1<=len(r)<=8 or any(not isinstance(v,(str,int,float,type(None))) or len(str(v))>200 for v in r) for r in table)):
            raise ValueError(f'{index}枚目の表は1〜20行・1〜8列、各セル200文字までです（1行目が見出し）。')
        diagram=slide.get('diagram')
        if diagram is not None and (not isinstance(diagram,dict) or set(diagram)-{'type','direction','nodes'} or diagram.get('type','flow') not in ('flow','stack')
                or diagram.get('direction','horizontal') not in ('horizontal','vertical') or not isinstance(diagram.get('nodes'),list)
                or not 2<=len(diagram['nodes'])<=8 or any(not isinstance(n,str) or not 1<=len(n)<=40 for n in diagram['nodes'])):
            raise ValueError(f'{index}枚目の図は type=flow|stack、nodes=2〜8個（各40文字まで）です。')
        duration=slide.get('duration',8)
        if type(duration) not in (int,float) or not 1<=duration<=180:raise ValueError('章の長さは1〜180秒です。')
    from team_document_capabilities import require_supported,runtime
    require_supported(path.suffix[1:]);python,ffmpeg=runtime()
    path.parent.mkdir(parents=True,exist_ok=True);path=output_path(root,args['path'])
    with tempfile.TemporaryDirectory(prefix='saikuru-output-',dir=path.parent) as folder:
        assets=[];render_args=dict(args,slides=[dict(s) for s in slides])
        render_args.pop('design_profile',None)
        if design:
            render_args['design']=dict(design_profile.slide_palette(design),body_file=design_profile.font_file(design,'body'),heading_file=design_profile.font_file(design,'heading'))
        for index,slide in enumerate(render_args['slides']):
            if not slide.get('asset_path'):continue
            name=Path(slide['asset_path'])
            if name.is_absolute() or '..' in name.parts:raise ValueError('素材は保存先内の相対パスを指定してください。')
            asset=(root/name).resolve(strict=True)
            if not asset.is_relative_to(root.resolve()) or not asset.is_file() or asset.suffix.lower() not in {'.png','.jpg','.jpeg','.mp4'}:raise ValueError('素材は保存先内のPNG・JPEG・MP4のみです。')
            if asset.stat().st_size>100*1024*1024:raise ValueError('素材は100MB以内です。')
            scene=slide.get('scene','実画面')
            if scene not in {'実画面','モデル設定','台帳','履歴'}:raise ValueError('素材のsceneは実画面・モデル設定・台帳・履歴です。')
            if len(slide.get('body',''))>140:raise ValueError('実画面を使う章の本文は140字以内にしてください。')
            copied=Path(folder)/('asset-'+str(index)+asset.suffix.lower());shutil.copyfile(asset,copied)
            slide['asset_path']=str(copied)
            assets.append({'path':str(asset.relative_to(root.resolve())),'sha256':hashlib.sha256(copied.read_bytes()).hexdigest(),'scene':scene,'slide':index+1})
        spec=Path(folder)/'input.json';generated=Path(folder)/('result'+path.suffix)
        spec.write_text(json.dumps(render_args,ensure_ascii=False),encoding='utf-8')
        result=subprocess.run([python,'-X','utf8',str(Path(__file__).with_name('document_media_worker.py')),str(spec),str(generated),ffmpeg or ''],stdout=subprocess.DEVNULL,stderr=subprocess.PIPE,timeout=1500,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        if result.returncode or not generated.is_file():
            detail=[l[len('SAIKURU_INPUT_ERROR: '):] for l in result.stderr.decode('utf-8','replace').splitlines() if l.startswith('SAIKURU_INPUT_ERROR: ')]
            if detail:raise ValueError('成果物の生成に失敗しました（入力を直せば作成できます）: '+detail[-1])
            raise ValueError('成果物の生成に失敗しました。VOICEVOX・話者・文字量・制作環境を確認してください。完成扱いにはしません。')
        if generated.stat().st_size>500*1024*1024:raise ValueError('成果物は500MB以内です。')
        previous=_set_aside(root,path) if replace else None  # only after generation succeeded
        with path.open('xb') as dest,generated.open('rb') as src:shutil.copyfileobj(src,dest)
    record={'path':str(path.relative_to(root)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size,'narration':all(bool(s.get('narration','').strip()) for s in slides),'visual_review':'未確認'}
    if previous:record['previous_version']=previous
    record['assets']=assets
    if job_id:
        manifest=root/('.saikuru-output-'+job_id+'.json')
        if manifest.is_symlink():raise ValueError('成果物管理ファイルにリンクは使えません。')
        records=json.loads(manifest.read_text(encoding='utf-8')) if manifest.exists() else []
        records.append(record);manifest.write_text(json.dumps(records,ensure_ascii=False),encoding='utf-8')
    return record

def office_text(path, limit=16000):
    """Text view of a Word/Excel file for reading or editing (no macros, nothing executed)."""
    _safe_office_package(path)
    lines = []
    if path.suffix == '.docx':
        from docx import Document
        doc = Document(str(path))
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(('# ' if p.style.name.startswith(('Heading', 'Title', '見出し')) else '') + p.text)
        for index, table in enumerate(doc.tables, 1):
            lines.append(f'[表{index}]')
            lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
    else:
        from openpyxl import load_workbook
        book = load_workbook(str(path), read_only=True, data_only=False)
        for sheet in book.worksheets:
            lines.append(f'[シート] {sheet.title}')
            for row in sheet.iter_rows(max_row=500, max_col=50):
                values = [(c.coordinate + '=' + str(c.value)) for c in row if c.value is not None]
                if values:
                    lines.append(', '.join(values))
        book.close()
    text = '\n'.join(lines)
    return {'content': text[:limit], 'truncated': len(text) > limit}


def pptx_text(path, limit=16000):
    """Text and design view of a PPTX for review: per-slide text in reading order, real tables, diagram
    boxes, notes, and a tally of fonts / text colours / fills / sizes. Nothing is rendered or executed."""
    _safe_office_package(path)
    from collections import Counter
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(str(path))
    fonts, colors, fills, sizes = Counter(), Counter(), Counter(), Counter()
    lines, slides = [], []

    def rgb(color):
        try:
            return '#' + str(color.rgb).lower() if color is not None and color.type is not None else None
        except (AttributeError, KeyError, ValueError, TypeError):
            return None

    def runs(frame):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                chars = len(run.text.strip())
                ea = run._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
                ea = ea.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea') if ea is not None else None
                name = (ea.get('typeface') if ea is not None else None) or run.font.name or paragraph.font.name
                if name:
                    fonts[name] += chars
                # Run settings first, then the paragraph defaults (saikuru's PPTX writer sets those).
                color = rgb(run.font.color) or rgb(paragraph.font.color)
                if color:
                    colors[color] += chars
                size = run.font.size or paragraph.font.size
                if size:
                    sizes[round(size.pt, 1)] += chars

    for index, slide in enumerate(prs.slides, 1):
        lines.append(f'[スライド{index}]')
        info = {'slide': index, 'tables': [], 'boxes': 0, 'pictures': 0}
        shapes = sorted(slide.shapes, key=lambda s: (int(s.top or 0) // Emu(300000), int(s.left or 0)))
        for shape in shapes:
            try:
                if shape.fill.type == 1:
                    fills['#' + str(shape.fill.fore_color.rgb).lower()] += 1
            except (AttributeError, TypeError, ValueError, KeyError, NotImplementedError):
                pass
            if shape.shape_type == 13:
                info['pictures'] += 1
                lines.append('[画像]')
            if getattr(shape, 'has_table', False) and shape.has_table:
                table = shape.table
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                info['tables'].append({'rows': len(rows), 'cols': len(table.columns)})
                lines.append(f'[表 {len(rows)}行×{len(table.columns)}列]')
                lines.extend(' | '.join(r) for r in rows)
                for row in table.rows:
                    for cell in row.cells:
                        runs(cell.text_frame)
                continue
            if getattr(shape, 'has_text_frame', False) and shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip()
                if shape.shape_type == 1 and getattr(shape, 'auto_shape_type', None) is not None and int(shape.auto_shape_type) != 1:
                    info['boxes'] += 1
                    lines.append('[図形] ' + text)
                else:
                    lines.extend(text.splitlines())
                runs(shape.text_frame)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append('[発表者メモ] ' + slide.notes_slide.notes_text_frame.text.strip())
        slides.append(info)
    text = '\n'.join(lines)
    return {'content': text[:limit], 'truncated': len(text) > limit, 'slide_count': len(prs.slides),
            'slide_size_in': [round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)],
            'slides': slides,
            'design': {'fonts': dict(fonts.most_common(6)), 'text_colors': dict(colors.most_common(8)),
                       'fills': dict(fills.most_common(8)), 'font_sizes_pt': dict(sizes.most_common(8))},
            'note': '文字・表・図形・書式は読み取り済み。見た目（配置の崩れ・はみ出し・読みやすさ）は画像での確認が別途必要です。'}


def _office_base(roots, name, suffix):
    """Existing Word/Excel to edit: in the output folder or the read-only source; copied, never modified."""
    relative = Path(name)
    if relative.is_absolute() or any(p in ('.', '..') or p.startswith('.') for p in relative.parts) or relative.suffix.lower() != suffix:
        raise ValueError('編集元は同じ形式の相対パスを指定してください。')
    for root in roots:
        candidate = (root / relative).resolve()
        if candidate.is_relative_to(root) and candidate.is_file() and not candidate.is_symlink():
            if candidate.stat().st_size > 50 * 1024 * 1024:
                raise ValueError('編集元は50MB以内です。')
            return candidate
    raise ValueError('編集元が見つかりません。')


def _cell_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if not isinstance(value, str) or len(value) > 5000:
        raise ValueError('セルの値は5000文字以内の文字・数値・真偽値です。')
    if value.startswith('=') and re.search(r'://|\[|[|\\]|\b(?:WEBSERVICE|FILTERXML|CALL|REGISTER|REGISTER\.ID|RTD|DDE|HYPERLINK|INDIRECT|IMAGE)\s*\(', value, re.I):
        raise ValueError('外部参照・外部呼び出しを含む数式は使えません。')
    return value


def _safe_office_package(path):
    """Reject active/external Office content before a document reader copies it."""
    if path.stat().st_size > 50*1024*1024:raise ValueError('Office資料は50MB以内です。')
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        if len(entries)>20000 or sum(e.file_size for e in entries)>200*1024*1024:
            raise ValueError('Office資料の展開サイズ・項目数が上限を超えています。')
        for entry in entries:
            name=entry.filename.lower()
            if entry.flag_bits & 1 or any(x in name for x in ('vbaproject','/embeddings/','/externallinks/','/activex/')):
                raise ValueError('マクロ・埋め込み・外部参照を持つOffice資料は編集できません。')
            if name.endswith('.xml') or name.endswith('.rels'):
                if entry.file_size>20000000:raise ValueError('Office XMLが大きすぎます。')
                raw=archive.read(entry)
                if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():raise ValueError('Office XMLの外部定義は使えません。')
            if name.endswith('.rels'):
                if entry.file_size>2000000:raise ValueError('Office関連情報が大きすぎます。')
                raw=archive.read(entry)
                if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():raise ValueError('Office XMLの外部定義は使えません。')
                for rel in ET.fromstring(raw):
                    if rel.get('TargetMode','').casefold()=='external':
                        raise ValueError('外部リンクを持つOffice資料は編集できません。')


def generate_office(roots, args, writable, job_id):
    """Create a new .docx/.xlsx, optionally starting from an existing file (base_path). Never overwrites."""
    if not writable:
        raise ValueError('制作はbuilderのみです。')
    root = roots[0]
    path = output_path(root, args.get('path', ''))
    if path.suffix not in ('.docx', '.xlsx'):
        raise ValueError('DOCX・XLSXの名前を指定してください。')
    replace = _own_output(root, path, job_id)
    from team_document_capabilities import require_supported
    require_supported(path.suffix[1:])
    base = _office_base(roots, args['base_path'], path.suffix) if args.get('base_path') else None
    if base:_safe_office_package(base)
    path.parent.mkdir(parents=True, exist_ok=True)
    path = output_path(root, args['path'])
    with tempfile.TemporaryDirectory(prefix='saikuru-office-', dir=path.parent) as folder:
        generated = Path(folder) / ('result' + path.suffix)
        if path.suffix == '.docx':
            from docx import Document
            doc = Document(str(base)) if base else Document()
            if args.get('design_profile'):
                # Styles always; page size/margins only for a new document (an edited base keeps its layout).
                design_profile.apply_docx(doc, _design(root, args['design_profile']), page_setup=not base)
            replaced = 0
            for item in (args.get('replacements') or [])[:200]:
                find, new = str(item.get('find', '')), str(item.get('replace', ''))
                if not find or len(find) > 1000 or len(new) > 5000:
                    raise ValueError('置換は1〜1000文字の検索文字列と5000文字以内の置換後文字列です。')
                cells = [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
                for p in list(doc.paragraphs) + cells:
                    if find in p.text:
                        text = p.text.replace(find, new)
                        for run in p.runs[1:]:
                            run.text = ''
                        if p.runs:
                            p.runs[0].text = text
                        else:
                            p.add_run(text)
                        replaced += 1
            if args.get('title'):
                doc.add_heading(str(args['title'])[:200], level=0)
            sections = args.get('sections') or []
            if len(sections) > 200:
                raise ValueError('節は200件までです。')
            for section in sections:
                if section.get('heading'):
                    doc.add_heading(str(section['heading'])[:200], level=min(max(int(section.get('level', 1)), 1), 3))
                for text in section.get('paragraphs') or []:
                    doc.add_paragraph(str(text)[:5000])
                for text in section.get('bullets') or []:
                    doc.add_paragraph(str(text)[:2000], style='List Bullet')
                table = section.get('table') or []
                if table:
                    if len(table) > 500 or any(len(r) > 30 for r in table):
                        raise ValueError('表は500行・30列までです。')
                    width = max(len(r) for r in table)
                    grid = doc.add_table(rows=len(table), cols=width)
                    grid.style = 'Table Grid'
                    for r, row in enumerate(table):
                        for c, value in enumerate(row):
                            grid.cell(r, c).text = '' if value is None else str(value)[:2000]
                    if args.get('design_profile'):
                        design_profile.style_docx_table(grid, _design(root, args['design_profile']))
            doc.save(str(generated))
            summary = {'replaced_paragraphs': replaced, 'added_sections': len(sections)}
        else:
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import Font
            book = load_workbook(str(base)) if base else Workbook()
            if base:
                for sheet in book.worksheets:
                    if sheet.max_row>5000 or sheet.max_column>100:raise ValueError('編集元は各シート5000行・100列以内です。')
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.data_type=='f':_cell_value(cell.value)
            if not base:
                book.remove(book.active)
            updates = args.get('updates') or []
            if len(updates) > 5000:
                raise ValueError('セル更新は5000件までです。')
            for item in updates:
                sheet = book[str(item['sheet'])] if str(item.get('sheet', '')) in book.sheetnames else None
                if sheet is None or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', str(item.get('cell', ''))):
                    raise ValueError('更新先のシート名とセル番地（例: B3）を指定してください。')
                sheet[item['cell']] = _cell_value(item.get('value'))
            sheets = args.get('sheets') or []
            if len(sheets) > 20:
                raise ValueError('シートは20件までです。')
            for spec in sheets:
                name = str(spec.get('name', ''))[:31]
                if not name or re.search(r'[\\/*?:\[\]]', name):
                    raise ValueError('シート名が不正です。')
                rows = spec.get('rows') or []
                if len(rows) > 5000 or any(len(r) > 100 for r in rows):
                    raise ValueError('行は5000行・100列までです。')
                existed = name in book.sheetnames
                sheet = book[name] if existed else book.create_sheet(name)
                # Rows for an existing, non-empty sheet are appended below its data.
                start = sheet.max_row + 1 if existed and (sheet.max_row > 1 or sheet['A1'].value is not None) else 1
                for r, row in enumerate(rows, start):
                    for c, value in enumerate(row, 1):
                        sheet.cell(r, c, _cell_value(value))
                if spec.get('header') and rows and start == 1:
                    for cell in sheet[1]:
                        cell.font = Font(bold=True)
                    sheet.freeze_panes = 'A2'
                for letter, width in (spec.get('column_widths') or {}).items():
                    if re.fullmatch(r'[A-Z]{1,3}', str(letter)) and isinstance(width, (int, float)) and 1 <= width <= 200:
                        sheet.column_dimensions[letter].width = width
            if not book.sheetnames:
                raise ValueError('シートを1つ以上指定してください。')
            book.save(str(generated))
            summary = {'updated_cells': len(updates), 'sheets': book.sheetnames}
        previous = _set_aside(root, path) if replace else None  # only after the new file was built
        with path.open('xb') as dest, generated.open('rb') as src:
            shutil.copyfileobj(src, dest)
    record = {'path': str(path.relative_to(root)), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
              'bytes': path.stat().st_size, 'edited_from': args.get('base_path') or None, 'assets': [],
              'visual_review': '未確認', 'summary': summary}
    if previous:
        record['previous_version'] = previous
    if job_id:
        manifest = root / ('.saikuru-output-' + job_id + '.json')
        if manifest.is_symlink():
            raise ValueError('成果物管理ファイルにリンクは使えません。')
        records = json.loads(manifest.read_text(encoding='utf-8')) if manifest.exists() else []
        records.append(record)
        manifest.write_text(json.dumps(records, ensure_ascii=False), encoding='utf-8')
    return record


def binary_documents(root, limit=200):
    """Generated PDF/PPTX/MP4/DOCX/XLSX in the output folder; list_files only shows text files."""
    found = []
    for directory, dirs, names in os.walk(root, followlinks=False):
        dirs[:] = [d for d in dirs if not d.startswith('.') and not (Path(directory) / d).is_symlink() and (Path(directory)/d).resolve().is_relative_to(root.resolve())]
        for name in names:
            path = Path(directory) / name
            if path.suffix.lower() in BINARY and not name.startswith('.') and not path.is_symlink() and path.resolve().is_relative_to(root.resolve()):
                found.append({'path': str(path.relative_to(root)), 'bytes': path.stat().st_size,
                              'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
                if len(found) >= limit:
                    return found
    return found


def main():
    source, output = (Path(p).resolve(strict=True) for p in sys.argv[1:3])
    writable = '--write' in sys.argv[3:]
    job_id=sys.argv[sys.argv.index('--job-id')+1] if '--job-id' in sys.argv else ''
    if job_id and not re.fullmatch('[a-f0-9]{32}',job_id):raise ValueError('Invalid job ID')
    # Output may be the source folder or inside it (writes are create-only document files);
    # an ancestor of the source would widen the write scope, so it stays refused.
    if output != source and source.is_relative_to(output):
        raise ValueError('保存先に読み取り元を含む上位フォルダは指定できません。')
    tools = [dict(t, description='List bounded non-sensitive text files of the read-only source. output_documents lists generated PDF/PPTX/MP4/DOCX/XLSX in the output folder with bytes and SHA256.') if t['name'] == 'list_files' else t for t in TOOLS] + [{'name': 'read_document', 'description': 'Read a document in the approved output folder. Word/Excel files are returned as text (headings, paragraphs, tables, cells) with their SHA256. PowerPoint (.pptx) returns per-slide text in reading order, tables ([表 r行×c列] + rows), diagram boxes ([図形]), speaker notes, slide_count, per-slide table/box/picture counts, and a design tally (fonts, text colours, fills, font sizes) to compare with a design profile. PDF returns page_count and per-page text (up to 60 pages / 16,000 characters). Visual layout still needs a rendered check.',
        'inputSchema': {'type':'object','properties':{'path':{'type':'string'}},'required':['path'],'additionalProperties':False}}]
    tools.append({'name':'media_environment','description':'Check local movie renderer and list available VOICEVOX speaker names/style IDs. No settings changes.', 'inputSchema':{'type':'object','properties':{},'additionalProperties':False}})
    if writable:
        tools.append({'name':'generate_media','description':'Generate editable PPTX, slide PDF, or MP4 from structured slides. PPTX is designed: theme colour, cover/section/content layouts, real tables (table: rows, first row = header), box-and-arrow diagrams (diagram: {type: flow|stack, direction, nodes}), page numbers and speaker notes. Use layout=cover for the first slide and section for chapter dividers; prefer tables/diagrams over long text. If a page does not fit, the error names the page; split it and retry. VOICEVOX narration uses the local engine and explicit speaker_id. Pages are always landscape 16:9 slides (PDF included); A4 or portrait pages are not possible, so make an A4 document with generate_office (docx) instead. A file this request generated can be re-generated under the same name (the previous version moves to .saikuru-versions/); any other existing file is never overwritten. Return paths and hashes; viewing/listening remains unverified.',
            'inputSchema':{'type':'object','properties':{'path':{'type':'string'},'speaker_id':{'type':'integer'},'design_profile':{'type':'string','description':'Relative *.design.json from learn_design; applies the learned colours and fonts'},'theme':{'type':'object','properties':{'accent':{'type':'string','description':'#RRGGBB'}},'additionalProperties':False},'document_title':{'type':'string'},'slides':{'type':'array','minItems':1,'maxItems':80,'items':{'type':'object','properties':{'title':{'type':'string'},'body':{'type':'string'},'layout':{'type':'string','enum':['cover','section','content']},'table':{'type':'array','maxItems':20,'items':{'type':'array','maxItems':8,'items':{'type':['string','number','null']}}},'diagram':{'type':'object','properties':{'type':{'type':'string','enum':['flow','stack']},'direction':{'type':'string','enum':['horizontal','vertical']},'nodes':{'type':'array','minItems':2,'maxItems':8,'items':{'type':'string'}}},'required':['nodes'],'additionalProperties':False},'notes':{'type':'string'},'narration':{'type':'string'},'duration':{'type':'number'},'asset_path':{'type':'string','description':'Relative PNG/JPEG/MP4 path in approved output folder; integrated in this slide'},'scene':{'type':'string','enum':['実画面','モデル設定','台帳','履歴']}},'required':['title','body'],'additionalProperties':False}}},'required':['path','slides'],'additionalProperties':False}})
        cell = {'type':['string','number','boolean','null']}
        tools.append({'name':'generate_office','description':'Create a new Word (.docx) or Excel (.xlsx) file in the approved output folder. To edit an existing file, set base_path (relative, in the output folder or read-only source); the result is saved under the new path and the original is never modified. Word: replacements, title, sections(heading/level/paragraphs/bullets/table). Excel: updates(sheet/cell/value), sheets(name/rows/header/column_widths; rows append to an existing sheet). Formulas allowed except external references.',
            'inputSchema':{'type':'object','properties':{'path':{'type':'string'},'base_path':{'type':'string'},'design_profile':{'type':'string','description':'Relative *.design.json from learn_design; Word styles (fonts, sizes, colours) and, for a new file, page size and margins'},
                'title':{'type':'string'},
                'replacements':{'type':'array','maxItems':200,'items':{'type':'object','properties':{'find':{'type':'string'},'replace':{'type':'string'}},'required':['find','replace'],'additionalProperties':False}},
                'sections':{'type':'array','maxItems':200,'items':{'type':'object','properties':{'heading':{'type':'string'},'level':{'type':'integer','minimum':1,'maximum':3},'paragraphs':{'type':'array','items':{'type':'string'}},'bullets':{'type':'array','items':{'type':'string'}},'table':{'type':'array','items':{'type':'array','items':cell}}},'additionalProperties':False}},
                'updates':{'type':'array','maxItems':5000,'items':{'type':'object','properties':{'sheet':{'type':'string'},'cell':{'type':'string'},'value':cell},'required':['sheet','cell','value'],'additionalProperties':False}},
                'sheets':{'type':'array','maxItems':20,'items':{'type':'object','properties':{'name':{'type':'string'},'rows':{'type':'array','items':{'type':'array','items':cell}},'header':{'type':'boolean'},'column_widths':{'type':'object','additionalProperties':{'type':'number'}}},'required':['name'],'additionalProperties':False}}},
                'required':['path'],'additionalProperties':False}})
        tools.append({'name':'write_document','description':'Create documents only in the approved output folder. Updating a generated document requires its SHA256.',
            'inputSchema':{'type':'object','properties':{'path':{'type':'string'},'content':{'type':'string'},'expected_sha256':{'type':'string'},'design_profile':{'type':'string','description':'HTML only: relative *.design.json; its CSS is inserted into <head>'}},'required':['path','content'],'additionalProperties':False}})
        tools.append({'name':'learn_design','description':'Learn a unified design from a reference PDF (in the read-only source or the output folder): colours (background, text, heading, primary, secondary, surface, muted), fonts mapped to installed fonts, type sizes, line height, page size and margins. Saves design/<name>.design.json and a Japanese summary design/<name>.design.md (never overwrites). Then pass design_profile to generate_media, generate_office (Word) or write_document (HTML). Reading only; nothing in the PDF is executed.',
            'inputSchema':{'type':'object','properties':{'pdf_path':{'type':'string','description':'Relative .pdf path'},'name':{'type':'string','description':'Profile name (letters, digits, Japanese, - _; max 40)'}},'required':['pdf_path'],'additionalProperties':False}})
    sys.stdin.reconfigure(encoding='utf-8');sys.stdout.reconfigure(encoding='utf-8')
    calls = 0
    for raw in sys.stdin:
        msg = {}
        try:
            if len(raw)>1500000: raise ValueError('Request too large')
            msg=json.loads(raw)
            if 'id' not in msg: continue
            method=msg.get('method'); params=msg.get('params',{})
            if method=='initialize': result={'protocolVersion':params.get('protocolVersion','2024-11-05'),'capabilities':{'tools':{}},'serverInfo':{'name':'document-scope','version':'1.0'}}
            elif method=='ping': result={}
            elif method=='tools/list': result={'tools':tools}
            elif method=='tools/call':
                calls+=1
                if calls>100: raise ValueError('Tool limit')
                name=params.get('name');args=params.get('arguments',{})
                value=(media_environment() if name=='media_environment' else generate_media(output,args,writable,job_id) if name=='generate_media' else generate_office((output,source),args,writable,job_id) if name=='generate_office' else learn_design((output,source),args,writable) if name=='learn_design' else document_operation(output,name,args,writable) if name in ('read_document','write_document') else dict(execute(source,name,args),output_documents=binary_documents(output)) if name=='list_files' else execute(source,name,args))
                result={'content':[{'type':'text','text':json.dumps(value,ensure_ascii=False)}]}
            else: raise ValueError('Unsupported operation')
            reply={'jsonrpc':'2.0','id':msg['id'],'result':result}
        except Exception as exc:
            reply={'jsonrpc':'2.0','id':msg.get('id'),'error':{'code':-32602,'message':str(exc)}}
        print(json.dumps(reply,ensure_ascii=False),flush=True)

if __name__=='__main__': main()
